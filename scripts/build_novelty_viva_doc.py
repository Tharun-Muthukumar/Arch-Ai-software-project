from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PROJECT_DOCUMENTATION.md"
OUTPUT = ROOT / "docs" / "ArchAI_Novelty_Accuracy_and_Viva_Guide.docx"

BLACK = "000000"
DARK_BLUE = "17365D"
MID_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "D9D9D9"
CODE_GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), LIGHT_GRAY)


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name: str = "Aptos", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_font(run)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Consolas", 9)
            run.font.color.rgb = RGBColor(45, 45, 45)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_font(run)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run)


def normalize_text(text: str) -> str:
    replacements = {
        "\u2192": "->",
        "\u2190": "<-",
        "\u2014": " - ",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2026": "...",
        "\u2265": ">=",
        "\u2264": "<=",
        "\u00d7": "x",
        "\u2248": "approximately ",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title._element.get_or_add_pPr()
    borders = title_p_pr.find(qn("w:pBdr"))
    if borders is not None:
        title_p_pr.remove(borders)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 18)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("ArchAI Novelty Accuracy and Viva Guide")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(14)
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run(
        "Code backed explanation of deterministic requirement engineering, "
        "accuracy safeguards, dynamic generation, bug corrections, and viva answers"
    )
    set_font(run, "Aptos", 14)
    run.font.color.rgb = RGBColor(50, 50, 50)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(18)
    intro.paragraph_format.space_after = Pt(10)
    add_inline(
        intro,
        "Purpose. This guide explains what is original in ArchAI beyond calling Ollama or exposing APIs. "
        "It maps every defensible claim to code, explains how functional and non-functional requirements "
        "are separated, states what is dynamic and what is deterministic, records the bugs corrected, "
        "and prepares the presenter for technical questioning.",
    )

    evidence = doc.add_paragraph()
    add_inline(
        evidence,
        "Verified state. 494 backend tests passed with Ollama disabled; 39 frontend tests passed; "
        "TypeScript compilation and the production build succeeded. The 44/44 classification result "
        "is a repository regression benchmark, not a claim of universal real-world accuracy.",
    )

    # Part II starts with its own page break in the Markdown renderer.
    heading = doc.add_paragraph("How to use this guide", style="Heading 1")
    heading.paragraph_format.space_after = Pt(8)
    for item in (
        "Read Parts II and III to understand the algorithms and novelty.",
        "Use the bug ledger to explain the engineering process and evidence.",
        "Memorise the short answers and numbers in Part IV before the viva.",
        "Demonstrate an unseen domain with Ollama disabled to prove dynamic behavior.",
        "Never claim 100 percent general accuracy; state the benchmark scope and limitations.",
    ):
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline(paragraph, item)
    doc.add_page_break()


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    mark_repeat_header(header)
    for index in range(columns):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(rows[0][index] if index < len(rows[0]) else "")
        set_font(run, size=9)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for index in range(columns):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, values[index] if index < len(values) else "")
            for run in paragraph.runs:
                set_font(run, run.font.name or "Aptos", 8.7)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_GRAY)
    p_pr.append(shading)
    run = paragraph.add_run("\n".join(lines).rstrip())
    set_font(run, "Consolas", 8.2)


def render_markdown(doc: Document, markdown: str) -> None:
    lines = normalize_text(markdown).splitlines()
    index = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if text:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, text)

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            raw_rows: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                values = [value.strip() for value in candidate.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    raw_rows.append(values)
                index += 1
            add_markdown_table(doc, raw_rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = re.sub(r"\*\*|`", "", heading.group(2)).strip()
            if text.startswith("Part "):
                doc.add_page_break()
                level = 1
            paragraph = doc.add_paragraph(text, style=f"Heading {min(level, 3)}")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue

        if stripped in {"---", "***"}:
            flush_paragraph()
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            item_parts = [(bullet or numbered).group(1)]
            index += 1
            while index < len(lines):
                continuation = lines[index].strip()
                if not continuation:
                    break
                if (
                    re.match(r"^(#{1,4})\s+", continuation)
                    or re.match(r"^[-*]\s+", continuation)
                    or re.match(r"^\d+\.\s+", continuation)
                    or continuation.startswith(("|", "```", ">"))
                    or continuation in {"---", "***"}
                ):
                    break
                item_parts.append(continuation)
                index += 1
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, " ".join(item_parts))
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
            continue

        if not stripped:
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    marker = "# Part II"
    if marker not in source:
        raise RuntimeError("Part II was not found in PROJECT_DOCUMENTATION.md")
    body = source[source.index(marker):]

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    doc.core_properties.title = "ArchAI Novelty Accuracy and Viva Guide"
    doc.core_properties.subject = "Technical explanation and viva preparation"
    doc.core_properties.author = "ArchAI Project Team"
    doc.core_properties.keywords = "ArchAI, requirements, architecture, deterministic, Ollama, viva"

    add_cover(doc)
    render_markdown(doc, body)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("ArchAI Novelty Accuracy and Viva Guide")
        set_font(run, size=8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)
        add_page_number(section.footer.paragraphs[0])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
